"""
app/poa/excel_writer.py
------------------------
Ported from the standalone RITM-POA tool's poa_excel.py. Writes the enriched
rows to an .xlsx workbook with the exact column order of the RITM-POA output
template, plus light styling (bold header, filters, frozen header row, auto
column width, and a highlight fill on any cell that still says
"MANUAL ENTRY REQUIRED" so it's easy to spot before sending it out).

Adapted from the original only in its I/O: `write_excel` now writes into an
in-memory `BytesIO` and returns the workbook's bytes, instead of saving to a
filesystem path — this module runs as part of a stateless API service with no
local file state (see ARCHITECTURE.md), and the generated workbook is stored
as bytes in `PoaRun.excel_content`.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.poa.enrich import OUTPUT_COLUMNS, MANUAL, INVALID_IP
from app.poa.lookup import SL_MSG_RULE_COUNT_FULL, SL_MSG_NOT_FILLED_SINGLE

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MANUAL_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
# Bad input is a different problem from a missing inventory row, so it gets
# its own fill (Excel's standard "bad" red) rather than the MANUAL yellow.
INVALID_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
EMPTY_OUTPUT_VALUE = "None"
# Every populated cell of this column is a validation failure, so the whole
# column is flagged rather than matching on the marker text alone.
VALIDATION_COLUMN = OUTPUT_COLUMNS.index("Input Validation") + 1


def _display_value(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return EMPTY_OUTPUT_VALUE
    return value


def write_excel(rows, sheet_name="Sheet1") -> bytes:
    """Build the enriched-rows workbook in memory and return its bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    ws.append(OUTPUT_COLUMNS)
    for col_idx, _ in enumerate(OUTPUT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for row_data in rows:
        ws.append([_display_value(row_data.get(col)) for col in OUTPUT_COLUMNS])

    # style data rows + auto width
    max_len = [len(c) for c in OUTPUT_COLUMNS]
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(OUTPUT_COLUMNS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.value == MANUAL:
                cell.fill = MANUAL_FILL
            elif (
                cell.value == INVALID_IP
                or cell.value in (SL_MSG_RULE_COUNT_FULL, SL_MSG_NOT_FILLED_SINGLE)
                or (c == VALIDATION_COLUMN and cell.value not in (None, EMPTY_OUTPUT_VALUE))
            ):
                cell.fill = INVALID_FILL
            val_len = len(str(cell.value)) if cell.value is not None else 0
            if val_len > max_len[c - 1]:
                max_len[c - 1] = val_len

    for i, width in enumerate(max_len, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 12), 45)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
